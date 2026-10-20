"""Export the Public Influence analysis and join its codes to the original graph.

Run from any directory: python scripts/export-public-influence.py
Use --check to verify that the committed export matches both source files.
Dependency: openpyxl (see scripts/requirements.txt).
"""

from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import sys

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
HEADERS = {
    "Node_Influence": (
        "Domain", "Frame12 Node Code", "Frame12 Node Title", "Public Action",
        "Participation Access", "Potential Influence Target",
        "Process Initiation / Convener", "Decision Authority", "Rights / Limitations",
    ),
    "Mechanism_Summary": (
        "Primary mechanism", "Node count", "Interpretation", "Codex visualization guidance",
    ),
    "Initiation_Exceptions": (
        "Node", "Mechanism", "Public-initiated action",
        "Eligibility / authority distinction", "Source treatment",
    ),
    "Codex_README": ("Field", "Instruction"),
}
MECHANISM_IDS = {
    "Comment / consultation / hearing / stakeholder input": "comment",
    "Mixed DPU participation": "formal",
    "Administrative / judicial challenge": "challenge",
    "Community negotiation": "negotiation",
    "Compliance complaint / reporting": "compliance",
}
DOMAIN_MECHANISMS = {
    "Boston Planning & Zoning": "comment",
    "Grid & Infrastructure": "comment",
    "Environmental Review": "comment",
    "Formal Participation": "formal",
    "Legal Challenges": "challenge",
    "Negotiation": "negotiation",
    "Compliance & Complaints": "compliance",
}


def require(condition, message):
    if not condition:
        raise ValueError(message)


def source_name(path):
    """Keep provenance portable and avoid publishing local filesystem paths."""
    try:
        return path.resolve().relative_to(ROOT).as_posix()
    except ValueError:
        return path.name


def export(workbook_path, graph_path):
    workbook = openpyxl.load_workbook(workbook_path, data_only=False)
    tables = {}
    hyperlinks = {}
    try:
        require(set(workbook.sheetnames) == set(HEADERS), "Workbook sheets changed; review the export schema.")
        for name, headers in HEADERS.items():
            sheet = workbook[name]
            rows = list(sheet.iter_rows())
            require(tuple(cell.value for cell in rows[0]) == headers, f"{name}: unexpected headers.")
            require(not sheet.merged_cells.ranges, f"{name}: merged cells need explicit handling.")
            tables[name] = []
            for row in rows[1:]:
                require(all(cell.data_type != "f" for cell in row), f"{name}: formulas need explicit handling.")
                values = [cell.value for cell in row]
                require(any(value is not None for value in values), f"{name}: unexpected empty row.")
                tables[name].append(dict(zip(headers, values)))
                for cell in row:
                    if cell.hyperlink:
                        hyperlinks[f"{name}!{cell.coordinate}"] = {
                            "target": cell.hyperlink.target,
                            "location": cell.hyperlink.location,
                            "tooltip": cell.hyperlink.tooltip,
                        }
    finally:
        workbook.close()

    # Universal-newline reading keeps this digest identical across Windows
    # checkouts (CRLF) and GitHub Actions checkouts (LF).
    graph_text = graph_path.read_text(encoding="utf-8")
    graph = json.loads(graph_text)
    graph_nodes = graph["nodes"]
    codes = [node["display_code"] for node in graph_nodes if node.get("display_code")]
    require(len(codes) == len(set(codes)), "Original graph has ambiguous display codes.")
    by_code = {node["display_code"]: node for node in graph_nodes if node.get("display_code")}
    source_codes = [row["Frame12 Node Code"] for row in tables["Node_Influence"]]
    require(len(source_codes) == len(set(source_codes)), "Node_Influence contains duplicate node codes.")

    def joined_node(code):
        require(code in by_code, f"Node code {code!r} does not exist in the original graph.")
        node = by_code[code]
        require(
            isinstance(node["node_id"], str) and re.fullmatch(r"\d{19}", node["node_id"]),
            f"{code}: original Miro ID must remain a 19-digit string.",
        )
        return node

    nodes = []
    for source_row, row in enumerate(tables["Node_Influence"], 2):
        code = row["Frame12 Node Code"]
        node = joined_node(code)
        domain = row["Domain"]
        require(domain in DOMAIN_MECHANISMS, f"{code}: unclassified domain {domain!r}.")
        require(node.get("Public Influence"), f"{code}: original graph has no Public Influence annotation.")
        nodes.append({
            "nodeId": node["node_id"],
            "code": code,
            "title": row["Frame12 Node Title"],
            "domain": domain,
            "mechanismId": DOMAIN_MECHANISMS[domain],
            "publicAction": row["Public Action"],
            "participationAccess": row["Participation Access"],
            "influenceTarget": row["Potential Influence Target"],
            "convener": row["Process Initiation / Convener"],
            "decisionAuthority": row["Decision Authority"],
            "rightsLimitations": row["Rights / Limitations"],
            "evidenceLinks": node.get("evidence_links", []),
            "originalPublicInfluence": node["Public Influence"],
            "sourceUrl": node.get("source_url"),
            "sourceRow": source_row,
        })

    original_annotated = {node["node_id"] for node in graph_nodes if node.get("Public Influence")}
    require(
        {node["nodeId"] for node in nodes} == original_annotated,
        "The analysis must cover exactly the original graph's Public Influence nodes.",
    )
    actual_counts = Counter(node["mechanismId"] for node in nodes)
    mechanisms = []
    total_rows = []
    for row in tables["Mechanism_Summary"]:
        label = row["Primary mechanism"]
        if label == "TOTAL":
            total_rows.append(row)
            continue
        require(label in MECHANISM_IDS, f"Unrecognized source mechanism: {label!r}.")
        mechanism_id = MECHANISM_IDS[label]
        require(actual_counts[mechanism_id] == row["Node count"], f"{label}: count does not match classified nodes.")
        mechanisms.append({
            "id": mechanism_id,
            "label": label,
            "count": row["Node count"],
            "interpretation": row["Interpretation"],
            "guidance": row["Codex visualization guidance"],
            "nodeCodes": [node["code"] for node in nodes if node["mechanismId"] == mechanism_id],
        })
    require(len(total_rows) == 1, "Mechanism_Summary must contain one total row.")
    require(len(mechanisms) == len(MECHANISM_IDS), "Mechanism_Summary must retain all five distinct categories.")
    require(len({item["id"] for item in mechanisms}) == len(mechanisms), "Duplicate mechanisms.")
    require(total_rows[0]["Node count"] == len(nodes), "Source total does not match annotated node count.")
    require(sum(item["count"] for item in mechanisms) == len(nodes), "Category counts do not cover the annotated nodes.")

    supplementary = []
    for source_row, row in enumerate(tables["Initiation_Exceptions"], 2):
        code = row["Node"]
        node = joined_node(code)
        supplementary.append({
            "code": code,
            "nodeId": node["node_id"],
            "mechanism": row["Mechanism"],
            "publicAction": row["Public-initiated action"],
            "eligibility": row["Eligibility / authority distinction"],
            "sourceTreatment": row["Source treatment"],
            "evidenceLinks": node.get("evidence_links", []),
            "sourceUrl": node.get("source_url"),
            "sourceRow": source_row,
        })

    guidance = [{"field": row["Field"], "instruction": row["Instruction"]} for row in tables["Codex_README"]]
    instructions = {row["field"]: row["instruction"] for row in guidance}
    return {
        "schemaVersion": 1,
        "source": {
            "workbook": source_name(workbook_path),
            "sha256": hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
            "graph": source_name(graph_path),
            "graphSha256": hashlib.sha256(graph_text.encode("utf-8")).hexdigest(),
            "classification": "Primary mechanisms follow the source workbook's domain grouping and summary counts.",
        },
        "summary": {
            "annotatedNodeCount": len(nodes),
            "supplementaryMechanismCount": len(supplementary),
            "interpretation": total_rows[0]["Interpretation"],
            "caveat": instructions["Caveat"],
        },
        "mechanisms": mechanisms,
        "nodes": nodes,
        "supplementary": supplementary,
        "guidance": guidance,
        "sourceTables": tables,
        "sourceHyperlinks": hyperlinks,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=ROOT / "database/Frame12_Public_Influence_Visualization_Database.xlsx")
    parser.add_argument("--graph", type=Path, default=ROOT / "web/data/frame12.json")
    parser.add_argument("--output", type=Path, default=ROOT / "web/data/public-influence.json")
    parser.add_argument("--check", action="store_true", help="Validate sources and compare the export without writing.")
    args = parser.parse_args()
    try:
        result = export(args.workbook, args.graph)
        text = json.dumps(result, ensure_ascii=False, indent=2) + "\n"
        if args.check:
            require(args.output.is_file(), "Public Influence export is missing; run the exporter first.")
            require(args.output.read_text(encoding="utf-8") == text, "Public Influence export is stale; run the exporter.")
        else:
            args.output.parent.mkdir(parents=True, exist_ok=True)
            args.output.write_text(text, encoding="utf-8")
        verb = "Verified" if args.check else "Exported"
        print(f"{verb} {len(result['nodes'])} annotated nodes, {len(result['mechanisms'])} mechanisms, "
              f"{len(result['supplementary'])} supplementary pathways; all graph IDs and evidence links preserved.")
    except (ValueError, KeyError, OSError) as error:
        print(f"Public Influence export failed: {error}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
